"""Builds the sample product sheet owners can download, fill in and upload."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HEAD = ["ID", "Name", "Price", "Stock", "Category", "Description", "Image"]
ROWS = [
    ["S001", "Black Canvas Sneakers", 2500, 12, "Shoes", "Comfortable everyday sneakers. Sizes 38-45.", ""],
    ["S002", "Leather Wallet (Brown)", 900, 30, "Accessories", "Genuine leather, 6 card slots.", ""],
    ["S003", "Kikoy Beach Towel", 1200, 8, "Home", "Soft cotton kikoy, 100 x 180 cm.", ""],
    ["S004", "Phone Case - Samsung A14", 450, 25, "Phone accessories", "Shockproof clear case.", ""],
    ["S005", "Wireless Earbuds", 3200, 5, "Electronics", "Bluetooth 5.3, 20 hours with charging case.", ""],
    ["S006", "Handwoven Basket (Large)", 1500, "", "Home", "Made in Kenya. Leave Stock empty if you don't want to count stock.", ""],
]


def build() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    font = Font(name="Arial", size=10)
    ws.append(HEAD)
    for r in ROWS:
        ws.append(r)
    for row in ws.iter_rows():
        for c in row:
            c.font = font
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for c in ws[1]:
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F7A4D")
    for col, w in zip("ABCDEFG", [8, 28, 10, 8, 20, 50, 36]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        for c in row:
            c.number_format = "#,##0"
    ws.freeze_panes = "A2"

    guide = wb.create_sheet("How to fill in")
    lines = [
        ("How to fill in the Products sheet", True),
        ("Delete the sample rows and type your own products, one product per row. Keep the first row (headings).", False),
        ("", False),
        ("ID - a short unique code for each product (e.g. S001). Optional, but recommended: keep it the same even if you rename the product.", False),
        ("Name - what customers will see. Required.", False),
        ("Price - number only (no 'KES'). Required.", False),
        ("Stock - how many you have. When something sells, the assistant reduces it automatically. When you type a new number here, it is used.", False),
        ("      Leave empty if you don't want to track stock (the item is then always available).", False),
        ("Category - e.g. Shoes. Helps customers browse. Optional.", False),
        ("Description - sizes, colours, materials. The assistant uses this to answer questions. Optional but very helpful.", False),
        ("Image - a direct link to a photo (a link ending in .jpg or .png). Optional.", False),
    ]
    for text, bold in lines:
        guide.append([text])
        guide.cell(guide.max_row, 1).font = Font(name="Arial", size=11 if bold else 10, bold=bold)
    guide.column_dimensions["A"].width = 130
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
