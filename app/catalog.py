"""Reads the product list from an uploaded Excel file or a Google Sheet link."""
from __future__ import annotations

import csv
import difflib
import io
import re
import time

import httpx
from openpyxl import load_workbook

from . import db, settings

XLSX_FILE = settings.DATA_DIR / "products.xlsx"

ALIASES = {
    "name": ["name", "item", "product", "product name", "title"],
    "price": ["price", "cost", "amount", "selling price"],
    "stock": ["stock", "qty", "quantity", "available"],
    "description": ["description", "details", "desc", "about"],
    "category": ["category", "type", "group"],
    "image": ["image", "image url", "image_url", "photo", "picture", "image link"],
    "id": ["id", "sku", "code"],
}

_cache = {"t": 0.0, "key": None, "items": [], "warnings": []}


def _sheet_csv_url(url: str) -> str:
    m = re.search(r"/spreadsheets/d/([a-zA-Z0-9_-]+)", url)
    if not m:
        return url
    out = f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=csv"
    gid = re.search(r"[#&?]gid=(\d+)", url)
    return out + (f"&gid={gid.group(1)}" if gid else "")


def _to_number(value):
    if value is None or str(value).strip() == "":
        return None
    cleaned = re.sub(r"[^\d.]", "", str(value).replace(",", ""))
    try:
        return float(cleaned) if cleaned else None
    except ValueError:
        return None


def _rows_from_xlsx(path) -> list[list]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Products"] if "Products" in wb.sheetnames else wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _rows_from_csv(text: str) -> list[list]:
    return [row for row in csv.reader(io.StringIO(text))]


def _parse(rows: list[list]) -> tuple[list[dict], list[str]]:
    warnings: list[str] = []
    # find the header row: first row that contains a recognisable name and price column
    header_idx, colmap = None, {}
    for i, row in enumerate(rows[:10]):
        cells = [str(c).strip().lower() if c is not None else "" for c in row]
        m = {}
        for key, names in ALIASES.items():
            for j, cell in enumerate(cells):
                if cell in names and key not in m:
                    m[key] = j
        if "name" in m and "price" in m:
            header_idx, colmap = i, m
            break
    if header_idx is None:
        return [], ["Could not find the header row. The first row needs at least 'Name' and 'Price' columns."]

    items, used_ids = [], set()
    for n, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        def cell(key):
            j = colmap.get(key)
            return row[j] if j is not None and j < len(row) else None

        name = str(cell("name") or "").strip()
        if not name or name.lower().startswith(("example:", "(example")):
            continue
        price = _to_number(cell("price"))
        if price is None:
            warnings.append(f"Row {n} ('{name}') skipped: price is missing or not a number.")
            continue
        stock_raw = _to_number(cell("stock"))
        pid = str(cell("id") or "").strip() or f"P{n}"
        if pid in used_ids:
            warnings.append(f"Row {n}: ID '{pid}' is used twice; renamed.")
            pid = f"{pid}-{n}"
        used_ids.add(pid)
        items.append(
            dict(
                id=pid,
                name=name,
                price=price,
                stock=int(stock_raw) if stock_raw is not None else None,
                description=str(cell("description") or "").strip(),
                category=str(cell("category") or "").strip(),
                image=str(cell("image") or "").strip(),
            )
        )
    return items, warnings


def load(force: bool = False) -> list[dict]:
    sheet_url = settings.get("sheet_url")
    key = (sheet_url, XLSX_FILE.stat().st_mtime if XLSX_FILE.exists() else None)
    if not force and key == _cache["key"] and time.time() - _cache["t"] < 60:
        return _cache["items"]

    rows, warnings = [], []
    try:
        if sheet_url:
            r = httpx.get(_sheet_csv_url(sheet_url), follow_redirects=True, timeout=20)
            r.raise_for_status()
            rows = _rows_from_csv(r.content.decode("utf-8-sig"))
        elif XLSX_FILE.exists():
            rows = _rows_from_xlsx(XLSX_FILE)
        else:
            warnings.append("No products yet. Upload an Excel file or paste a Google Sheet link.")
    except Exception as e:  # keep serving the last good list if the sheet is temporarily unreachable
        warnings.append(f"Could not read the product list: {e}")
        if _cache["items"]:
            _cache["warnings"] = warnings
            return _cache["items"]

    items, parse_warnings = _parse(rows) if rows else ([], [])
    _cache.update(t=time.time(), key=key, items=items, warnings=warnings + parse_warnings)
    db.sync_inventory(items)
    return items


def warnings() -> list[str]:
    return _cache["warnings"]


def with_stock(p: dict) -> dict:
    """Product plus the live stock number (None = not tracked)."""
    live = db.stock_of(p["id"]) if p["stock"] is not None else None
    return {**p, "stock": live if p["stock"] is not None else None}


def find(pid: str) -> dict | None:
    for p in load():
        if p["id"].lower() == str(pid).strip().lower():
            return with_stock(p)
    return None


def search(query: str = "", limit: int = 8) -> list[dict]:
    items = load()
    tokens = [t for t in re.findall(r"\w+", (query or "").lower()) if len(t) > 1]
    if not tokens:
        return [with_stock(p) for p in items[:limit]]
    scored = []
    for p in items:
        name, cat, desc = p["name"].lower(), p["category"].lower(), p["description"].lower()
        words = re.findall(r"\w+", name)
        score = 0.0
        for t in tokens:
            if t in name:
                score += 3
            if t in cat:
                score += 2
            if t in desc:
                score += 1
            if difflib.get_close_matches(t, words, n=1, cutoff=0.8):
                score += 2
        if score:
            scored.append((score, p))
    scored.sort(key=lambda x: -x[0])
    return [with_stock(p) for _, p in scored[:limit]]


def categories() -> list[str]:
    return sorted({p["category"] for p in load() if p["category"]})
